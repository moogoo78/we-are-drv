"""端對端 API 測試：以暫存 SQLite 與上傳目錄執行。

執行：docker run --rm -v $PWD/api:/app -w /app <含 pytest 的映像> python -m pytest -q tests
"""
import io
import os
import tempfile

import pytest

_tmp = tempfile.mkdtemp()
os.environ.setdefault("JWT_SECRET", "test-secret")
os.environ["DATABASE_PATH"] = os.path.join(_tmp, "test.sqlite3")
os.environ["UPLOAD_DIR"] = os.path.join(_tmp, "uploads")
os.environ.pop("SMTP_HOST", None)

from app import create_app  # noqa: E402

PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
       b"\x00\x00\x00\rIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82")


@pytest.fixture(scope="module")
def client():
    app = create_app()
    app.testing = True
    return app.test_client()


@pytest.fixture(scope="module")
def admin(client):
    res = client.post("/api/auth/login", json={"username": "admin", "password": "changeme123"})
    return {"Authorization": "Bearer " + res.get_json()["token"]}


def open_batch_id(client):
    return next(b["id"] for b in client.get("/api/batches").get_json()["batches"] if b["status"] == "open")


def register(client, id_number, **extra):
    data = {
        "batch_id": str(open_batch_id(client)), "name": "王小明", "id_number": id_number, "birth_date": "1990-05-12",
        "sex": "男", "phone": "0912-345-678", "email": "a@example.com", "district": "安樂區", "village": "樂一里",
        "address": "基隆市安樂區樂一里1號", "meal_pref": "葷食", "support_pref": "願意", "agree_consent": "true",
        "specialty": '["醫護或急救相關背景"]', "photo": (io.BytesIO(PNG), "p.png", "image/png"),
    }
    data.update(extra)
    return client.post("/api/registrations", data=data, content_type="multipart/form-data")


def test_public_endpoints(client):
    assert len(client.get("/api/public/villages").get_json()["districts"]["安樂區"]) == 25
    assert client.get("/api/public/sessions").status_code == 200
    assert client.get("/api/public/instructors").get_json()["instructors"]
    assert "material_key" not in client.get("/api/batches").get_json()["batches"][0]


def test_registration_rejects_wrong_village(client):
    res = register(client, "B123456789", village="仁壽里")
    assert res.status_code == 400


def test_full_training_flow(client, admin):
    res = register(client, "A123456789")
    assert res.status_code == 201, res.get_json()
    reg_no = res.get_json()["reg_no"]
    reg = client.get(f"/api/registrations?q={reg_no}", headers=admin).get_json()["registrations"][0]
    assert reg["village"] == "樂一里"

    # 審核 → 正取
    assert client.patch(f"/api/registrations/{reg['id']}", json={"status": "approved"}, headers=admin).status_code == 200

    # 學員登入（手機格式不同仍可比對）
    res = client.post("/api/me/login", json={"id_number": "a123456789", "birth_date": "1990-05-12", "phone": "0912345678"})
    assert res.status_code == 200
    student = {"Authorization": "Bearer " + res.get_json()["token"]}
    assert client.post("/api/me/login", json={"id_number": "A123456789", "birth_date": "1990-05-12",
                                              "phone": "0900000000"}).status_code == 401
    # 學員 token 不可呼叫後台 API
    assert client.get("/api/registrations", headers=student).status_code == 403
    assert client.get("/api/me/overview", headers=admin).status_code == 403

    # 題庫 + 開放線上測驗
    client.post("/api/admin/quiz/questions", json={"qtype": "tf", "question": "Q1", "answer": "O"}, headers=admin)
    client.post("/api/admin/quiz/questions", json={"qtype": "choice", "question": "Q2", "options": ["a", "b", "c", "d"],
                                                   "answer": "D"}, headers=admin)
    batch_id = reg["batch_id"]
    client.patch(f"/api/batches/{batch_id}", json={"quiz_open": True, "survey_open": True}, headers=admin)

    ov = client.get("/api/me/overview", headers=student).get_json()
    mine = ov["registrations"][0]
    assert mine["quiz_available"] and mine["survey_available"] and mine["material_key"]
    quiz = client.get(f"/api/me/quiz/{reg['id']}", headers=student).get_json()
    assert all("answer" not in q for q in quiz["questions"])
    answers = {str(q["id"]): ("O" if q["qtype"] == "tf" else "D") for q in quiz["questions"]}
    assert client.post(f"/api/me/quiz/{reg['id']}", json={"answers": answers}, headers=student).get_json()["passed"]
    assert client.post(f"/api/me/quiz/{reg['id']}", json={"answers": answers}, headers=student).status_code == 409

    # 問卷
    qs = client.get("/api/me/survey").get_json()["questions"]
    ans = {}
    for q in qs:
        ans[q["id"]] = {"single": lambda q: q["options"][0], "multi": lambda q: [q["options"][0]],
                        "rating": lambda q: 5, "text": lambda q: ""}[q["type"]](q)
    assert client.post(f"/api/me/survey/{reg['id']}", json={"answers": ans}, headers=student).status_code == 201

    # 報到、術科
    client.patch(f"/api/registrations/{reg['id']}/training", json={"checked_in": True, "practical_passed": True},
                 headers=admin)

    # 證照：申請中 → 已核發 → 列入防災士名冊
    cert = client.post("/api/admin/certs", json={"reg_no": reg_no}, headers=admin).get_json()["certificate"]
    assert cert["apply_status"] == "applying"
    client.patch(f"/api/admin/certs/{cert['id']}", json={"apply_status": "issued"}, headers=admin)
    members = client.get("/api/admin/members?q=王小明", headers=admin).get_json()
    assert members["total"] == 1 and members["members"][0]["support_willing"] == 1
    stats = client.get("/api/admin/members/stats", headers=admin).get_json()
    assert stats["by_village"][0]["village"] == "樂一里"
    assert len(client.get("/api/admin/certs/unclaimed", headers=admin).get_json()["certificates"]) == 1

    # 學員申請補發、參訓證明
    for kind in ("reissue", "attendance_proof"):
        assert client.post("/api/me/applications", json={"registration_id": reg["id"], "kind": kind},
                           headers=student).status_code == 201
    apps = client.get("/api/admin/applications", headers=admin).get_json()["applications"]
    assert len(apps) == 2
    client.patch(f"/api/admin/applications/{apps[0]['id']}", json={"status": "done"}, headers=admin)

    # 通知（未設定 SMTP → 待發送）
    rec = client.get(f"/api/admin/notify/recipients?kind=approved&batch_id={batch_id}", headers=admin).get_json()
    ids = [r["registration_id"] for r in rec["recipients"]]
    res = client.post("/api/admin/notify/send", json={"channel": "sms", "category": "approved", "body": "{name} 正取",
                                                      "registration_ids": ids}, headers=admin).get_json()
    assert res["queued"] == len(ids)
    csv_text = client.get("/api/admin/notify/export.csv?channel=sms", headers=admin).get_data(as_text=True)
    assert "王小明 正取" in csv_text
    # 證照已核發未領取 → 證照領取提醒名單
    rec = client.get("/api/admin/notify/recipients?kind=cert_pickup_reminder", headers=admin).get_json()["recipients"]
    assert [r["registration_id"] for r in rec] == [reg["id"]]
    templates = client.get("/api/admin/notify/config", headers=admin).get_json()["templates"]
    for kind in ("course_reminder_3d", "needs_more_info_reminder", "cert_pickup_reminder", "cert_mailed"):
        assert kind in templates
        assert client.get(f"/api/admin/notify/recipients?kind={kind}", headers=admin).status_code == 200

    # 報表
    for path in ("registrations.xlsx", "roster.xlsx", "written.xlsx", "consent.xlsx", "survey.xlsx",
                 "roster-photos.zip", "practical.html", "consent.html", "summary.html", "registrations.html"):
        res = client.get(f"/api/admin/reports/{path}?batch_id={batch_id}", headers=admin)
        assert res.status_code == 200, path
    assert "王小明" in client.get(f"/api/admin/reports/registrations.html?batch_id={batch_id}",
                                  headers=admin).get_data(as_text=True)


def test_materials_and_photos(client, admin):
    batches = client.get("/api/batches/admin", headers=admin).get_json()["batches"]
    b = batches[0]
    client.post("/api/admin/courses/materials", data={"title": "講義", "visibility": "public",
                                                      "file": (io.BytesIO(b"%PDF-1.4"), "a.pdf")},
                headers=admin, content_type="multipart/form-data")
    client.post("/api/admin/courses/materials", data={"title": "學員教材", "visibility": "enrolled", "batch_id": b["id"],
                                                      "file": (io.BytesIO(b"%PDF-1.4"), "b.pdf")},
                headers=admin, content_type="multipart/form-data")
    pub = client.get("/api/public/materials").get_json()
    assert len(pub["public"]) == 1 and pub["enrolled"] == []
    enrolled = client.get(f"/api/public/materials?batch_id={b['id']}&key={b['material_key']}").get_json()["enrolled"]
    assert len(enrolled) == 1
    assert client.get(f"/api/public/materials/{enrolled[0]['id']}/download").status_code == 403
    assert client.get(f"/api/public/materials/{enrolled[0]['id']}/download?key={b['material_key']}").status_code == 200

    res = client.post("/api/admin/courses/photos", data={"batch_id": b["id"], "photo": (io.BytesIO(PNG), "x.png", "image/png")},
                      headers=admin, content_type="multipart/form-data")
    assert res.status_code == 201
    photo = client.get("/api/public/photos").get_json()["photos"][0]
    assert client.get(f"/api/public/photos/{photo['id']}").status_code == 200


def test_member_import_official_template(client, admin):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["", "", "", "", "完訓日期(發證日期)", "", "", "", "", "出生日期"])
    ws.merge_cells("E1:G1")
    ws.merge_cells("J1:L1")
    ws.append(["*照片編號", "*證書編號", "*姓名", "*培訓計畫名稱", "*年(西元)", "*月", "*日", "*培訓單位", "*身分證字號",
               "*年(西元)", "*月", "*日", "*性別", "市內電話", "*行動電話", "*E-mail", "*所屬縣市", "*居住地址",
               "最高學歷", "服務單位", "職稱", "*參與防災工作意願"])
    ws.append([1, "C-001", "陳大文", "112年第1期", 2023, 5, 20, "基隆市消防局", "C123456789", 1970, 1, 2, "男", "",
               "0911111111", "c@example.com", "基隆市", "基隆市信義區孝忠里3號", "", "里辦公處", "", "有"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    res = client.post("/api/admin/members/import", data={"file": (buf, "hist.xlsx")}, headers=admin,
                      content_type="multipart/form-data").get_json()
    assert res["inserted"] == 1, res
    m = client.get("/api/admin/members?q=陳大文", headers=admin).get_json()["members"][0]
    assert (m["birth_date"], m["trained_on"], m["district"], m["village"], m["support_willing"]) == \
        ("1970-01-02", "2023-05-20", "信義區", "孝忠里", 1)


def test_announcements(client, admin):
    client.post("/api/admin/announcements", json={"title": "未來公告", "published_at": "2999-01-01"}, headers=admin)
    client.post("/api/admin/announcements", json={"title": "今日公告", "category": "change"}, headers=admin)
    titles = [a["title"] for a in client.get("/api/public/announcements").get_json()["announcements"]]
    assert titles == ["今日公告"]
