"""成果報告一鍵產出：Excel 名冊／成績／問卷，及可列印（另存 PDF）的測驗表單、同意書、成果報告。"""
import base64
import io
import json
import os
import zipfile
from datetime import date
from html import escape

from flask import Blueprint, Response, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from ..auth import require_auth
from ..db import get_db
from ..files import UPLOAD_DIR, upload_path
from .. import survey_def

reports_bp = Blueprint("reports", __name__)

TRAINING_UNIT = "基隆市政府消防局"
ROSTER_COURSES = [
    "基礎急救訓練", "急救措施實作(含急救術科測驗)", "防災士職責與任務、我國災防體系與運作",
    "我國近年災害經驗及災害特性", "資訊掌握、運用與社區防災計畫", "個人與居家防護措施",
    "個人與居家防護措施（情境練習）", "社區防災工作推動與運作", "社區避難收容場所開設與運作", "防災計畫實作與驗證",
]
CONSENT_TEXT = (
    "基隆市政府消防局為辦理防災士培訓、證照核發、災時人力盤點及成果統計，依個人資料保護法第 8 條規定告知："
    "本表所蒐集之個人資料，僅於前述特定目的及法令規定範圍內使用，保存期限依檔案管理規定辦理。"
    "您得依法請求查詢、閱覽、複製、補充、更正、停止蒐集處理利用或刪除；未提供必要欄位者，將無法完成報名。"
)
STATUS_LABEL = {"pending": "審核中", "approved": "正取", "waitlist": "備取", "needs_more_info": "待補件", "rejected": "未錄取"}


def _batch(conn):
    batch_id = request.args.get("batch_id")
    if not batch_id:
        return None
    return conn.execute("SELECT * FROM batches WHERE id = ?", (batch_id,)).fetchone()


def _regs(conn, batch_id, approved_only=True):
    return conn.execute(
        f"""SELECT r.*, c.cert_no FROM registrations r
            LEFT JOIN certificates c ON c.registration_id = r.id
            WHERE r.batch_id = ? {"AND r.status = 'approved'" if approved_only else "AND r.status != 'rejected'"}
            ORDER BY r.id""",
        (batch_id,),
    ).fetchall()


def _xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _ymd(s):
    try:
        y, m, d = (int(x) for x in (s or "").split("-")[:3])
        return y, m, d
    except ValueError:
        return "", "", ""


def _html_page(title, body):
    return Response(f"""<!DOCTYPE html><html lang="zh-Hant"><head><meta charset="UTF-8"><title>{escape(title)}</title>
<style>
  body {{ font-family: "Noto Sans TC", "Microsoft JhengHei", sans-serif; color: #222; margin: 24px; font-size: 14px; }}
  h1 {{ font-size: 20px; text-align: center; margin: 0 0 4px; }}
  h2 {{ font-size: 16px; margin: 24px 0 8px; border-bottom: 1px solid #999; padding-bottom: 4px; }}
  .sub {{ text-align: center; color: #555; margin-bottom: 16px; }}
  table {{ border-collapse: collapse; width: 100%; margin-bottom: 12px; }}
  th, td {{ border: 1px solid #666; padding: 6px 8px; text-align: left; vertical-align: top; }}
  th {{ background: #eee; }}
  .page {{ page-break-after: always; }}
  .page:last-child {{ page-break-after: auto; }}
  .sign img {{ max-height: 60px; }}
  .box {{ display: inline-block; width: 14px; height: 14px; border: 1px solid #000; text-align: center; line-height: 14px; }}
  .toolbar {{ margin-bottom: 16px; }}
  @media print {{ .toolbar {{ display: none; }} body {{ margin: 0; }} }}
</style></head><body>
<div class="toolbar"><button onclick="window.print()">列印／另存 PDF</button></div>
{body}
</body></html>""", mimetype="text/html")


# ---- Excel ----

@reports_bp.route("/registrations.xlsx", methods=["GET"])
@require_auth
def registrations_xlsx():
    """報名名冊（基隆市防災士培訓報名表範本格式）。"""
    conn = get_db()
    batch = _batch(conn)
    if not batch:
        conn.close()
        return jsonify({"error": "請選擇梯次"}), 400
    rows = _regs(conn, batch["id"], approved_only=request.args.get("scope") != "all")
    conn.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "報名名冊"
    ws.append([f"基隆市防災士培訓 {batch['name']} 報名表"])
    ws.merge_cells("A1:V1")
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["編號", "姓名", "身分證字號", "出生西元年", "出生民國年", "出生月份", "出生日期", "性別", "戶籍行政區／里",
               "居住地址", "住家電話", "行動電話", "EMAIL", "服務單位", "職稱", "最高學歷", "緊急聯絡人姓名",
               "緊急聯絡人之關係", "緊急聯絡人之連絡電話", "飲食需求", "審核狀態", "報名序號"])
    _style_header(ws, 2)
    for i, r in enumerate(rows, start=1):
        y, m, d = _ymd(r["birth_date"])
        ws.append([i, r["name"], r["id_number"], y, (y - 1911) if y else "", m, d, r["sex"],
                   f'{r["district"]}{r["village"] or ""}', r["address"], r["home_phone"], r["phone"], r["email"],
                   r["unit_village"], r["job_title"], r["education"], r["emergency_name"], r["emergency_relation"],
                   r["emergency_phone"], r["meal_pref"], STATUS_LABEL.get(r["status"], r["status"]), r["reg_no"]])
    return _xlsx_response(wb, f"報名名冊_{batch['name']}.xlsx")


@reports_bp.route("/roster.xlsx", methods=["GET"])
@require_auth
def roster_xlsx():
    """防災士公版表格（學員名冊），供函報中央。"""
    conn = get_db()
    batch = _batch(conn)
    if not batch:
        conn.close()
        return jsonify({"error": "請選擇梯次"}), 400
    rows = _regs(conn, batch["id"])
    conn.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "防災士匯入公版表格"
    ws.append(["", "＊標示之欄位為必填", "", "", "完訓日期(發證日期)", "", "", "", "", "出生日期", "", "", "", "", "", "大寫",
               "", "", "", "", "", "", "", "合格/不合格", "合格/不合格", "參訓情形(已完成課程或抵免者打「●」)"])
    ws.merge_cells("E1:G1")
    ws.merge_cells("J1:L1")
    ws.merge_cells("Z1:AI1")
    ws.append(["*照片編號", "*證書編號", "*姓名", "*培訓計畫名稱", "*年(西元)", "*月", "*日", "*培訓單位", "*身分證字號",
               "*年(西元)", "*月", "*日", "*性別", "市內電話(xx-xxxxxxxx)", "*行動電話(09xx-xxxxxx)", "*E-mail",
               "*所屬縣市", "*居住地址", "最高學歷", "服務單位", "職稱", "*參與防災工作意願", "*學科測驗成績",
               "*術科測驗成績", "*認證結果是否合格", *ROSTER_COURSES])
    _style_header(ws, 1)
    _style_header(ws, 2)
    ey, em, ed = _ymd(batch["end_date"])
    for i, r in enumerate(rows, start=1):
        by, bm, bd = _ymd(r["birth_date"])
        practical = {1: "合格", 0: "不合格"}.get(r["practical_passed"], "")
        certified = "合格" if r["written_passed"] == 1 and r["practical_passed"] == 1 else (
            "不合格" if 0 in (r["written_passed"], r["practical_passed"]) else "")
        mark = "●" if r["checked_in"] else "○"
        ws.append([i, r["cert_no"] or "", r["name"], batch["name"], ey, em, ed, TRAINING_UNIT, r["id_number"],
                   by, bm, bd, r["sex"], r["home_phone"] or "", r["phone"], r["email"], "基隆市", r["address"],
                   r["education"] or "", r["unit_village"] or "", r["job_title"] or "",
                   "有" if r["support_pref"] in ("願意", "可支援") else "無",
                   r["written_score"] if r["written_score"] is not None else "", practical, certified,
                   *([mark] * len(ROSTER_COURSES))])
    return _xlsx_response(wb, f"防災士公版表格_學員名冊_{batch['name']}.xlsx")


@reports_bp.route("/roster-photos.zip", methods=["GET"])
@require_auth
def roster_photos():
    """學員證件照，以公版名冊的照片編號命名。"""
    conn = get_db()
    batch = _batch(conn)
    if not batch:
        conn.close()
        return jsonify({"error": "請選擇梯次"}), 400
    rows = _regs(conn, batch["id"])
    conn.close()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for i, r in enumerate(rows, start=1):
            path = os.path.join(UPLOAD_DIR, r["photo_path"] or "")
            if r["photo_path"] and os.path.isfile(path):
                z.write(path, f"{i}_{r['name']}{os.path.splitext(path)[1]}")
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"證件照_{batch['name']}.zip", mimetype="application/zip")


@reports_bp.route("/written.xlsx", methods=["GET"])
@require_auth
def written_xlsx():
    """學科測驗卷結果。"""
    conn = get_db()
    batch = _batch(conn)
    if not batch:
        conn.close()
        return jsonify({"error": "請選擇梯次"}), 400
    rows = conn.execute(
        """SELECT r.reg_no, r.name, r.written_score, r.written_passed, q.submitted_at, q.id AS attempt_id
           FROM registrations r LEFT JOIN quiz_attempts q ON q.registration_id = r.id
           WHERE r.batch_id = ? AND r.status = 'approved' ORDER BY r.id""",
        (batch["id"],),
    ).fetchall()
    conn.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "學科測驗結果"
    ws.append(["編號", "報名序號", "姓名", "成績", "結果", "作答方式", "作答時間"])
    _style_header(ws)
    for i, r in enumerate(rows, start=1):
        result = {1: "合格", 0: "不合格"}.get(r["written_passed"], "未測驗")
        ws.append([i, r["reg_no"], r["name"], r["written_score"] if r["written_score"] is not None else "", result,
                   "線上測驗" if r["attempt_id"] else ("人工登錄" if r["written_score"] is not None else ""),
                   r["submitted_at"] or ""])
    scored = [r["written_score"] for r in rows if r["written_score"] is not None]
    ws.append([])
    ws.append(["", "", "應試人數", len(scored)])
    ws.append(["", "", "平均分數", round(sum(scored) / len(scored), 1) if scored else ""])
    ws.append(["", "", "合格人數", sum(1 for r in rows if r["written_passed"] == 1)])
    return _xlsx_response(wb, f"學科測驗結果_{batch['name']}.xlsx")


@reports_bp.route("/consent.xlsx", methods=["GET"])
@require_auth
def consent_xlsx():
    """個人資料同意書總表。"""
    conn = get_db()
    batch = _batch(conn)
    if not batch:
        conn.close()
        return jsonify({"error": "請選擇梯次"}), 400
    rows = _regs(conn, batch["id"], approved_only=request.args.get("scope") != "all")
    conn.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "個資同意書總表"
    ws.append(["編號", "報名序號", "姓名", "身分證字號", "同意方式", "同意時間", "是否同意"])
    _style_header(ws)
    for i, r in enumerate(rows, start=1):
        ws.append([i, r["reg_no"], r["name"], r["id_number"], "線上報名勾選", r["created_at"],
                   "同意" if r["agree_consent"] else "未同意"])
    return _xlsx_response(wb, f"個資同意書總表_{batch['name']}.xlsx")


def _survey_rows(conn, batch_id):
    return conn.execute(
        """SELECT s.answers, s.submitted_at, r.name, r.id_number, r.sex, r.id AS rid
           FROM survey_responses s JOIN registrations r ON r.id = s.registration_id
           WHERE r.batch_id = ? ORDER BY s.id""",
        (batch_id,),
    ).fetchall()


def _survey_averages(rows):
    sums, counts = {}, {}
    for r in rows:
        answers = json.loads(r["answers"])
        for q in survey_def.QUESTIONS:
            v = answers.get(q["id"])
            if q["type"] == "rating" and isinstance(v, int):
                sums[q["id"]] = sums.get(q["id"], 0) + v
                counts[q["id"]] = counts.get(q["id"], 0) + 1
    return {k: round(sums[k] / counts[k], 2) for k in sums}


@reports_bp.route("/survey.xlsx", methods=["GET"])
@require_auth
def survey_xlsx():
    """課後問卷調查（公版欄位順序）＋ 各題平均。"""
    conn = get_db()
    batch = _batch(conn)
    if not batch:
        conn.close()
        return jsonify({"error": "請選擇梯次"}), 400
    rows = _survey_rows(conn, batch["id"])
    conn.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "問卷回覆"
    ws.append(["主辦單位", "結訓日期", "一、學員編號", "二、姓名", "三、身分證字號", "四、性別",
               *[q["label"] for q in survey_def.QUESTIONS]])
    _style_header(ws)
    for i, r in enumerate(rows, start=1):
        a = json.loads(r["answers"])
        values = []
        for q in survey_def.QUESTIONS:
            v = a.get(q["id"], "")
            values.append(", ".join(v) if isinstance(v, list) else v)
        ws.append([TRAINING_UNIT, batch["end_date"], i, r["name"], r["id_number"], r["sex"], *values])
    avg = _survey_averages(rows)
    ws2 = wb.create_sheet("滿意度統計")
    ws2.append(["題目", "平均分數（1～5）"])
    _style_header(ws2)
    for q in survey_def.QUESTIONS:
        if q["type"] == "rating":
            ws2.append([q["label"], avg.get(q["id"], "")])
    ws2.append([])
    ws2.append(["填答人數", len(rows)])
    return _xlsx_response(wb, f"課後問卷調查_{batch['name']}.xlsx")


# ---- 可列印頁面 ----

@reports_bp.route("/practical.html", methods=["GET"])
@require_auth
def practical_html():
    """術科測驗表單：帶入學員姓名與教官簽名檔，勾選合格。"""
    conn = get_db()
    batch = _batch(conn)
    if not batch:
        conn.close()
        return jsonify({"error": "請選擇梯次"}), 400
    rows = _regs(conn, batch["id"])
    instructor = None
    if request.args.get("instructor_id"):
        instructor = conn.execute("SELECT * FROM instructors WHERE id = ?", (request.args["instructor_id"],)).fetchone()
    conn.close()
    sign = escape(instructor["name"]) if instructor else ""
    if instructor and instructor["signature_path"]:
        path = upload_path("signatures", instructor["signature_path"])
        if os.path.isfile(path):
            mime = "image/png" if path.endswith(".png") else "image/jpeg"
            with open(path, "rb") as f:
                sign = f'<img src="data:{mime};base64,{base64.b64encode(f.read()).decode()}" alt="{escape(instructor["name"])} 簽名">'
    trs = "".join(
        f"""<tr><td>{i}</td><td>{escape(r["name"])}</td><td>{escape(r["reg_no"] or "")}</td>
            <td style="text-align:center"><span class="box">{"✔" if r["practical_passed"] == 1 else ""}</span> 合格
            　<span class="box">{"✔" if r["practical_passed"] == 0 else ""}</span> 不合格</td>
            <td class="sign">{sign}</td></tr>"""
        for i, r in enumerate(rows, start=1)
    )
    body = f"""<h1>{escape(TRAINING_UNIT)} 防災士培訓 術科測驗（CPR／AED）紀錄表</h1>
<div class="sub">{escape(batch["name"])}　測驗日期：{escape(batch["end_date"])}　地點：{escape(batch["location"])}</div>
<table><thead><tr><th style="width:48px">編號</th><th>學員姓名</th><th>報名序號</th><th style="width:200px">測驗結果</th><th style="width:180px">測驗教官簽名</th></tr></thead>
<tbody>{trs or '<tr><td colspan="5">尚無錄取學員</td></tr>'}</tbody></table>"""
    return _html_page(f"術科測驗紀錄表_{batch['name']}", body)


@reports_bp.route("/consent.html", methods=["GET"])
@require_auth
def consent_html():
    """個人資料蒐集告知暨同意書（每位學員一頁）。"""
    conn = get_db()
    batch = _batch(conn)
    if not batch:
        conn.close()
        return jsonify({"error": "請選擇梯次"}), 400
    rows = _regs(conn, batch["id"], approved_only=request.args.get("scope") != "all")
    reg_id = request.args.get("registration_id")
    if reg_id:
        rows = [r for r in rows if str(r["id"]) == reg_id]
    conn.close()
    pages = "".join(
        f"""<div class="page"><h1>個人資料蒐集告知暨同意書</h1><div class="sub">{escape(batch["name"])}</div>
<p style="line-height:1.9">{escape(CONSENT_TEXT)}</p>
<table><tr><th style="width:140px">姓名</th><td>{escape(r["name"])}</td></tr>
<tr><th>身分證字號</th><td>{escape(r["id_number"])}</td></tr>
<tr><th>報名序號</th><td>{escape(r["reg_no"] or "")}</td></tr>
<tr><th>同意方式</th><td>{"☑ 已於線上報名時勾選同意" if r["agree_consent"] else "☐ 未同意"}</td></tr>
<tr><th>同意時間</th><td>{escape(r["created_at"])}</td></tr>
<tr><th>本人簽名</th><td style="height:60px"></td></tr></table></div>"""
        for r in rows
    )
    return _html_page(f"個資同意書_{batch['name']}", pages or "<p>尚無資料</p>")


@reports_bp.route("/registrations.html", methods=["GET"])
@require_auth
def registrations_html():
    """報名名冊（可列印／另存 PDF）。"""
    conn = get_db()
    batch = _batch(conn)
    if not batch:
        conn.close()
        return jsonify({"error": "請選擇梯次"}), 400
    rows = _regs(conn, batch["id"], approved_only=request.args.get("scope") != "all")
    conn.close()
    trs = "".join(
        f"""<tr><td>{i}</td><td>{escape(r["name"])}</td><td>{escape(r["id_number"])}</td><td>{escape(r["birth_date"])}</td>
            <td>{escape(r["sex"])}</td><td>{escape(r["district"])}{escape(r["village"] or "")}</td><td>{escape(r["phone"])}</td>
            <td>{escape(r["email"])}</td><td>{escape(r["unit_village"] or "")}</td><td>{escape(r["job_title"] or "")}</td>
            <td>{escape(r["emergency_name"] or "")} {escape(r["emergency_phone"] or "")}</td>
            <td>{escape(r["meal_pref"])}</td><td>{escape(STATUS_LABEL.get(r["status"], r["status"]))}</td></tr>"""
        for i, r in enumerate(rows, start=1)
    )
    body = f"""<style>@page {{ size: A4 landscape; }} table {{ font-size: 12px; }}</style>
<h1>基隆市防災士培訓 {escape(batch["name"])} 報名名冊</h1>
<div class="sub">{escape(batch["start_date"])}～{escape(batch["end_date"])}　{escape(batch["location"])}　共 {len(rows)} 人</div>
<table><thead><tr><th>編號</th><th>姓名</th><th>身分證字號</th><th>出生日期</th><th>性別</th><th>行政區／里</th><th>行動電話</th>
<th>Email</th><th>服務單位</th><th>職稱</th><th>緊急聯絡人</th><th>飲食</th><th>審核狀態</th></tr></thead>
<tbody>{trs or '<tr><td colspan="13">尚無資料</td></tr>'}</tbody></table>
<p style="color:#555">產出日期：{date.today().isoformat()}</p>"""
    return _html_page(f"報名名冊_{batch['name']}", body)


@reports_bp.route("/summary.html", methods=["GET"])
@require_auth
def summary_html():
    """各梯次成果報告。"""
    conn = get_db()
    batch = _batch(conn)
    if not batch:
        conn.close()
        return jsonify({"error": "請選擇梯次"}), 400
    bid = batch["id"]
    regs = conn.execute("SELECT * FROM registrations WHERE batch_id = ?", (bid,)).fetchall()
    sessions = conn.execute(
        """SELECT cs.*, i.name AS instructor_name FROM course_sessions cs
           LEFT JOIN instructors i ON i.id = cs.instructor_id WHERE cs.batch_id = ?
           ORDER BY cs.session_date, cs.start_time""",
        (bid,),
    ).fetchall()
    certs = conn.execute(
        """SELECT c.apply_status, c.pickup_status FROM certificates c
           JOIN registrations r ON r.id = c.registration_id WHERE r.batch_id = ?""",
        (bid,),
    ).fetchall()
    photos = conn.execute("SELECT id FROM course_photos WHERE batch_id = ?", (bid,)).fetchall()
    survey_rows = _survey_rows(conn, bid)
    conn.close()

    approved = [r for r in regs if r["status"] == "approved"]
    checked = [r for r in approved if r["checked_in"]]
    written = [r for r in approved if r["written_passed"] == 1]
    practical = [r for r in approved if r["practical_passed"] == 1]
    completed = [r for r in approved if r["checked_in"] and r["written_passed"] == 1 and r["practical_passed"] == 1]

    def pct(a, b):
        return f"{round(len(a) * 100 / len(b))}%" if b else "—"

    def dist(key):
        out = {}
        for r in approved:
            out[r[key] or "未填寫"] = out.get(r[key] or "未填寫", 0) + 1
        return "、".join(f"{escape(k)} {v} 人" for k, v in sorted(out.items(), key=lambda x: -x[1])) or "—"

    avg = _survey_averages(survey_rows)
    ratings = [v for k, v in avg.items() if k != "overall"]
    stats = [
        ("報名人數", len(regs)), ("正取人數", len(approved)),
        ("備取人數", sum(1 for r in regs if r["status"] == "waitlist")),
        ("報到人數（報到率）", f"{len(checked)}（{pct(checked, approved)}）"),
        ("學科合格（合格率）", f"{len(written)}（{pct(written, approved)}）"),
        ("術科合格（合格率）", f"{len(practical)}（{pct(practical, approved)}）"),
        ("結訓人數（結訓率）", f"{len(completed)}（{pct(completed, approved)}）"),
        ("證照已核發", sum(1 for c in certs if c["apply_status"] == "issued")),
        ("證照已領取", sum(1 for c in certs if c["pickup_status"] == "collected")),
        ("課後問卷填答", f"{len(survey_rows)} 份"),
        ("課程滿意度平均", f"{round(sum(ratings) / len(ratings), 2)} / 5" if ratings else "—"),
        ("整體規劃滿意度", f'{avg["overall"]} / 5' if "overall" in avg else "—"),
        ("課程照片", f"{len(photos)} 張"),
    ]
    stat_rows = "".join(f"<tr><th style='width:220px'>{escape(k)}</th><td>{escape(str(v))}</td></tr>" for k, v in stats)
    session_rows = "".join(
        f"<tr><td>{escape(s['session_date'])} {escape(s['start_time'] or '')}</td><td>{escape(s['topic'])}</td>"
        f"<td>{escape(s['instructor_name'] or '—')}</td><td>{s['hours']}</td></tr>"
        for s in sessions
    ) or "<tr><td colspan='4'>尚無課程資料</td></tr>"
    body = f"""<h1>基隆市防災士培訓 成果報告</h1>
<div class="sub">{escape(batch["name"])}　{escape(batch["start_date"])}～{escape(batch["end_date"])}　{escape(batch["location"])}</div>
<h2>一、辦理成果</h2><table>{stat_rows}</table>
<h2>二、課程內容</h2><table><thead><tr><th>時間</th><th>課程</th><th>講師</th><th>時數</th></tr></thead><tbody>{session_rows}</tbody></table>
<h2>三、學員組成（正取）</h2><table>
<tr><th style='width:220px'>性別</th><td>{dist("sex")}</td></tr>
<tr><th>行政區</th><td>{dist("district")}</td></tr>
<tr><th>職業別</th><td>{dist("occupation")}</td></tr></table>
<p style="color:#555">產出日期：{date.today().isoformat()}</p>"""
    return _html_page(f"成果報告_{batch['name']}", body)
