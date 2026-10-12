"""防災士管理：防災士名冊（歷年匯入＋本系統結訓）、歷史資料查詢、防災士統計。"""
import csv
import io
import json
import re
import sqlite3

from flask import Blueprint, g, jsonify, request
from openpyxl import load_workbook

from ..auth import require_auth, require_role
from ..db import get_db
from ..geo import VILLAGES, parse_district_village
from ..log import log_action
from ..members import json_list
from .dashboard import age_bucket

members_bp = Blueprint("members_admin", __name__)

ID_PATTERN = re.compile(r"^[A-Z][0-9A-Z][0-9]{8}$")


def _norm(h):
    return re.sub(r"[\s*＊()（）/]", "", str(h or ""))


def _cell(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _ymd(y, m, d):
    y, m, d = _cell(y), _cell(m), _cell(d)
    if not (y and m and d):
        return None
    try:
        year = int(y)
        if year < 1911:
            year += 1911  # 民國年
        return f"{year:04d}-{int(m):02d}-{int(d):02d}"
    except ValueError:
        return None


def _rows_from_upload(file):
    """讀取 xlsx 或 csv，回傳 [{正規化欄名: 值}]。支援消防署「防災士公版表格」兩列表頭格式。"""
    name = (file.filename or "").lower()
    if name.endswith(".csv"):
        text = file.read().decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        table = [row for row in reader]
    else:
        wb = load_workbook(file, read_only=False, data_only=True)
        ws = wb.worksheets[0]
        # 合併儲存格的群組標題（例如「出生日期」橫跨 年/月/日 三欄）往右展開
        group = {}
        for rng in ws.merged_cells.ranges:
            v = ws.cell(rng.min_row, rng.min_col).value
            for c in range(rng.min_col, rng.max_col + 1):
                group[(rng.min_row, c)] = v
        table = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            table.append([
                (group.get((r_idx, c_idx)) if v is None and (r_idx, c_idx) in group else v)
                for c_idx, v in enumerate(row, start=1)
            ])

    header_idx = next((i for i, row in enumerate(table[:10]) if any(_norm(c) == "姓名" for c in row)), None)
    if header_idx is None:
        raise ValueError("找不到「姓名」欄位，請確認檔案格式")
    header = [_norm(c) for c in table[header_idx]]
    above = [_norm(c) for c in table[header_idx - 1]] if header_idx > 0 else [""] * len(header)
    keys = []
    for i, h in enumerate(header):
        prefix = above[i] if i < len(above) else ""
        keys.append(prefix + h if h in ("年西元", "年", "月", "日") else h)
    rows = []
    for row in table[header_idx + 1:]:
        rec = {keys[i]: row[i] for i in range(min(len(keys), len(row))) if keys[i]}
        if _cell(rec.get("姓名")):
            rows.append(rec)
    return rows


def _pick(rec, *names):
    for n in names:
        for k, v in rec.items():
            if k.startswith(n) and _cell(v):
                return _cell(v)
    return ""


def _member_from_record(rec):
    id_number = _pick(rec, "身分證字號").upper()
    birth = _pick(rec, "出生日期") if not any(k.startswith("出生日期年") for k in rec) else ""
    if not birth:
        birth = _ymd(_pick(rec, "出生日期年"), _pick(rec, "出生日期月"), _pick(rec, "出生日期日"))
    trained = _pick(rec, "完訓日期")
    if any(k.startswith("完訓日期發證日期年") or k.startswith("完訓日期年") for k in rec):
        trained = _ymd(_pick(rec, "完訓日期發證日期年", "完訓日期年"), _pick(rec, "完訓日期發證日期月", "完訓日期月"),
                       _pick(rec, "完訓日期發證日期日", "完訓日期日"))
    address = _pick(rec, "居住地址", "通訊地址", "地址")
    district = _pick(rec, "行政區")
    village = _pick(rec, "里別", "里")
    if not district or not village:
        d, v = parse_district_village(address)
        district = district or d
        village = village or v
    specialty = _pick(rec, "專長")
    willing = _pick(rec, "參與防災工作意願", "是否可支援", "災時支援")
    return {
        "name": _pick(rec, "姓名"),
        "id_number": id_number or None,
        "sex": _pick(rec, "性別"),
        "birth_date": birth or None,
        "phone": _pick(rec, "行動電話", "手機", "電話"),
        "email": _pick(rec, "E-mail", "Email", "EMAIL", "電子郵件"),
        "address": address,
        "district": district or None,
        "village": village or None,
        "occupation": _pick(rec, "職業"),
        "unit": _pick(rec, "服務單位", "所屬單位"),
        "specialty": json.dumps([s.strip() for s in re.split(r"[,，、]", specialty) if s.strip()], ensure_ascii=False),
        "support_willing": 1 if willing in ("願意", "有", "是", "可支援", "Y", "y", "1") else 0,
        "cert_no": _pick(rec, "證書編號", "證照編號"),
        "training_name": _pick(rec, "培訓計畫名稱", "期別", "梯次"),
        "trained_on": trained or None,
    }


@members_bp.route("", methods=["GET"])
@require_auth
def list_members():
    conditions, params = [], []
    for arg, col in (("district", "district"), ("village", "village"), ("source", "source")):
        if request.args.get(arg):
            conditions.append(f"{col} = ?")
            params.append(request.args[arg])
    if request.args.get("support") == "1":
        conditions.append("support_willing = 1")
    if request.args.get("year"):
        conditions.append("substr(trained_on, 1, 4) = ?")
        params.append(request.args["year"])
    q = request.args.get("q")
    if q:
        conditions.append("(name LIKE ? OR id_number LIKE ? OR cert_no LIKE ? OR training_name LIKE ?)")
        params += [f"%{q}%"] * 4
    where = ("WHERE " + " AND ".join(conditions)) if conditions else ""
    page = max(int(request.args.get("page", 1) or 1), 1)
    page_size = min(int(request.args.get("page_size", 50) or 50), 500)

    conn = get_db()
    rows = conn.execute(
        f"""SELECT * FROM members {where} ORDER BY trained_on DESC, id DESC LIMIT ? OFFSET ?""",
        (*params, page_size, (page - 1) * page_size),
    ).fetchall()
    total = conn.execute(f"SELECT COUNT(*) AS n FROM members {where}", params).fetchone()["n"]
    conn.close()
    members = []
    for r in rows:
        m = dict(r)
        m["specialty"] = json_list(m["specialty"])
        members.append(m)
    return jsonify({"members": members, "total": total, "page": page, "page_size": page_size})


@members_bp.route("/stats", methods=["GET"])
@require_auth
def stats():
    conn = get_db()
    rows = conn.execute(
        "SELECT sex, birth_date, district, village, occupation, unit, specialty, support_willing, source FROM members"
    ).fetchall()
    conn.close()

    def count(values):
        out = {}
        for v in values:
            out[v] = out.get(v, 0) + 1
        return sorted(({"label": k, "n": n} for k, n in out.items()), key=lambda x: -x["n"])

    by_village = {}
    for r in rows:
        if r["district"] and r["village"]:
            key = (r["district"], r["village"])
            by_village[key] = by_village.get(key, 0) + 1
    ages = []
    for r in rows:
        try:
            ages.append(age_bucket(r["birth_date"]) if r["birth_date"] else "未填寫")
        except ValueError:
            ages.append("未填寫")
    specialties = [s for r in rows for s in json_list(r["specialty"])]
    return jsonify({
        "total": len(rows),
        "from_training": sum(1 for r in rows if r["source"] == "training"),
        "from_import": sum(1 for r in rows if r["source"] == "import"),
        "support_willing": sum(1 for r in rows if r["support_willing"]),
        "by_district": count(r["district"] or "未填寫" for r in rows),
        "by_village": sorted(
            ({"district": d, "village": v, "n": n} for (d, v), n in by_village.items()), key=lambda x: -x["n"]
        ),
        "villages_without_members": sum(
            1 for d, vs in VILLAGES.items() for v in vs if (d, v) not in by_village
        ),
        "by_sex": count(r["sex"] or "未填寫" for r in rows),
        "by_age": count(ages),
        "by_occupation": count(r["occupation"] or "未填寫" for r in rows),
        "by_unit": count(r["unit"] or "未填寫" for r in rows)[:20],
        "by_specialty": count(specialties),
    })


@members_bp.route("/import", methods=["POST"])
@require_auth
@require_role("admin", "staff")
def import_members():
    """匯入歷年防災士名單（消防署公版表格 xlsx 或 csv）；身分證字號相同者更新資料。"""
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "請選擇檔案"}), 400
    try:
        records = _rows_from_upload(file)
    except Exception as e:  # noqa: BLE001 — 檔案格式錯誤時回報給使用者
        return jsonify({"error": f"檔案讀取失敗：{e}"}), 400

    inserted = updated = 0
    errors = []
    conn = get_db()
    for i, rec in enumerate(records, start=1):
        m = _member_from_record(rec)
        if m["id_number"] and not ID_PATTERN.match(m["id_number"]):
            errors.append(f"第 {i} 筆（{m['name']}）身分證字號格式錯誤")
            continue
        existing = conn.execute("SELECT id, source FROM members WHERE id_number = ?", (m["id_number"],)).fetchone() \
            if m["id_number"] else None
        if existing:
            # 不以匯入資料覆寫本系統結訓者的資料
            if existing["source"] == "training":
                continue
            set_clause = ", ".join(f"{k} = ?" for k in m)
            conn.execute(f"UPDATE members SET {set_clause} WHERE id = ?", (*m.values(), existing["id"]))
            updated += 1
        else:
            try:
                conn.execute(
                    f"INSERT INTO members ({', '.join(m)}, source) VALUES ({', '.join('?' for _ in m)}, 'import')",
                    tuple(m.values()),
                )
                inserted += 1
            except sqlite3.IntegrityError as e:
                errors.append(f"第 {i} 筆（{m['name']}）：{e}")
    conn.commit()
    conn.close()
    log_action(g.user["name"], "匯入防災士名冊", f"{file.filename}：新增 {inserted}、更新 {updated}")
    return jsonify({"inserted": inserted, "updated": updated, "errors": errors[:50], "total": len(records)})


@members_bp.route("/<int:member_id>", methods=["DELETE"])
@require_auth
@require_role("admin")
def delete_member(member_id):
    conn = get_db()
    row = conn.execute("SELECT name FROM members WHERE id = ?", (member_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "找不到資料"}), 404
    conn.execute("DELETE FROM members WHERE id = ?", (member_id,))
    conn.commit()
    conn.close()
    log_action(g.user["name"], "刪除防災士資料", row["name"])
    return "", 204
